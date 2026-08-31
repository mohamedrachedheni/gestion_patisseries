"""
Commande : python manage.py import_donnees_excel --input CHEMIN [--noinput]

Recharge les tables de référence (voir core/management/reference_tables.py)
depuis un classeur Excel généré par export_donnees_excel — une feuille par
table, traitées dans le même ordre de dépendances de clé étrangère.

Les identifiants (colonne "id") sont préservés tels qu'exportés : une ligne
existante est mise à jour, une ligne absente est créée avec ce même id —
indispensable car des tables non couvertes ici (ventes, bons de livraison...)
référencent ces id par clé étrangère.

Opération non destructive (aucune suppression), mais peut écraser des
données existantes portant les mêmes id : à utiliser en connaissance de
cause, jamais sans sauvegarde préalable de la base cible.
"""
import datetime
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from core.management.reference_tables import nom_feuille, tables_reference


class Command(BaseCommand):
    help = 'Recharge les tables de référence depuis un classeur Excel (voir export_donnees_excel)'

    def add_arguments(self, parser):
        parser.add_argument('--input', required=True, help='Chemin du fichier .xlsx à recharger')
        parser.add_argument(
            '--noinput', action='store_true',
            help='Ne pas demander de confirmation avant d\'écrire en base',
        )

    def handle(self, *args, **options):
        chemin = options['input']
        try:
            wb = load_workbook(chemin, read_only=True, data_only=True)
        except FileNotFoundError:
            raise CommandError(f'Fichier introuvable : {chemin}')

        if not options['noinput']:
            reponse = input(
                f"Ceci va créer/mettre à jour des lignes dans la base courante à partir de "
                f"'{chemin}'. Sauvegardez la base cible avant de continuer si ce n'est pas déjà "
                f"fait. Continuer ? [oui/non] "
            )
            if reponse.strip().lower() not in ('oui', 'o', 'yes', 'y'):
                self.stdout.write('Annulé.')
                return

        self.stdout.write(self.style.MIGRATE_HEADING('\n=== Rechargement des tables de référence ==='))
        for model in tables_reference():
            feuille = nom_feuille(model)
            if feuille not in wb.sheetnames:
                self.stdout.write(self.style.WARNING(f'  • {model._meta.db_table} → feuille absente, ignorée'))
                continue

            ws = wb[feuille]
            lignes = ws.iter_rows(values_only=True)
            entetes = next(lignes)
            champs = {f.attname: f for f in model._meta.concrete_fields}

            compteur = 0
            with transaction.atomic():
                for brute in lignes:
                    if brute[0] is None:
                        continue
                    valeurs = {
                        nom: self._depuis_excel(champs[nom], val)
                        for nom, val in zip(entetes, brute)
                    }
                    pk = valeurs.pop('id')
                    model.objects.update_or_create(pk=pk, defaults=valeurs)
                    compteur += 1

            self.stdout.write(f'  • {model._meta.db_table} → {compteur} lignes rechargées')

        self.stdout.write(self.style.SUCCESS('\nRechargement terminé.\n'))

    @staticmethod
    def _depuis_excel(field, valeur):
        if valeur is None:
            return None
        interne = field.get_internal_type()
        if interne in (
            'AutoField', 'BigAutoField', 'IntegerField', 'SmallIntegerField',
            'PositiveIntegerField', 'PositiveSmallIntegerField', 'ForeignKey',
        ):
            return int(valeur)
        if interne == 'DecimalField':
            try:
                return Decimal(str(valeur))
            except InvalidOperation:
                return Decimal(0)
        if interne == 'BooleanField':
            return bool(valeur)
        if interne == 'DateField' and isinstance(valeur, datetime.datetime):
            # Colonne Date pure lue comme datetime par openpyxl (Excel ne
            # distingue pas les deux formats) : on retire la partie heure.
            return valeur.date()
        return valeur
